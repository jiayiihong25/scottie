"""PowerPoint extraction via python-pptx. One slide -> one SourceUnit."""

from __future__ import annotations

from pathlib import Path

from ..errors import EmptyExtractionError, ExtractionError
from ..models import Document, SourceUnit


def _shape_text(shape) -> list[str]:
    lines: list[str] = []
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            line = "".join(run.text for run in para.runs).strip()
            if line:
                lines.append(line)
    elif getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return lines


def _notes_text(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    frame = slide.notes_slide.notes_text_frame
    return (frame.text or "").strip() if frame else ""


def extract(path: Path) -> Document:
    from pptx import Presentation  # lazy: keep PDF-only runs light

    units: list[SourceUnit] = []
    try:
        prs = Presentation(str(path))
        for number, slide in enumerate(prs.slides, start=1):
            lines: list[str] = []
            heading = None
            title_shape = slide.shapes.title
            if title_shape is not None and title_shape.has_text_frame:
                title = title_shape.text_frame.text.strip()
                if title:
                    heading = title
                    lines.append(title)

            for shape in slide.shapes:
                if shape is title_shape:
                    continue
                lines.extend(_shape_text(shape))

            notes = _notes_text(slide)
            if notes:
                lines.append(f"[speaker notes] {notes}")

            if not lines:
                continue
            units.append(
                SourceUnit(number=number, text="\n".join(lines), heading=heading)
            )
    except Exception as exc:
        raise ExtractionError(f"could not read slide deck {path}: {exc}") from exc

    if not units:
        raise EmptyExtractionError(
            f"{path} produced no text — every slide is empty or image-only"
        )
    return Document(path=path, source_type="slides", units=units, unit_kind="slide")
